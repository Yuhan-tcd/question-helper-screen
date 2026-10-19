"""Archive course files and extract searchable text for the screenshot helper."""

import argparse
import hashlib
import re
import shutil
import subprocess
import tempfile
from pathlib import Path


LIBRARY = Path(__file__).resolve().parent / "course_materials"
SUPPORTED = {".pdf", ".ppt", ".pptx", ".docx", ".txt", ".md"}


def safe_name(value: str) -> str:
    name = re.sub(r"[^\w\-. ]+", "_", value, flags=re.UNICODE).strip(" ._")
    return name[:100] or "material"


def file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()[:10]


def read_plain_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeError:
            continue
    raise ValueError(f"无法读取文本编码：{path.name}")


def pdf_text(path: Path) -> tuple[str, list[str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    sections = []
    blank_pages = 0
    for number, page in enumerate(reader.pages, 1):
        body = (page.extract_text(extraction_mode="layout") or "").strip()
        if body:
            sections.append(f"## Page {number}\n\n{body}")
        else:
            blank_pages += 1
    warnings = []
    if blank_pages:
        warnings.append(f"{blank_pages} 页没有可提取文字；图片或扫描页需要另做识别")
    return "\n\n".join(sections), warnings


def pptx_text(path: Path) -> tuple[str, list[str]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(path))
    sections = []
    slides_without_text = 0
    slides_with_images = 0
    for number, slide in enumerate(presentation.slides, 1):
        parts = []
        has_image = False
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                parts.append(shape.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            if shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}:
                has_image = True
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes and notes.text.strip():
                parts.append("Speaker notes:\n" + notes.text.strip())
        if parts:
            sections.append(f"## Slide {number}\n\n" + "\n\n".join(parts))
        else:
            slides_without_text += 1
        slides_with_images += has_image
    warnings = []
    if slides_without_text:
        warnings.append(f"{slides_without_text} 张幻灯片没有可提取文字")
    if slides_with_images:
        warnings.append(f"{slides_with_images} 张幻灯片包含图片；图中的文字、公式和图表尚未识别")
    return "\n\n".join(sections), warnings


def docx_text(path: Path) -> tuple[str, list[str]]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    document = Document(str(path))
    parts = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    warnings = []
    if document.inline_shapes:
        warnings.append(f"Word 文件含 {len(document.inline_shapes)} 张嵌入图片；图片中的内容尚未识别")
    return "\n\n".join(parts), warnings


def convert_legacy_ppt(path: Path, output_dir: Path) -> Path:
    office = shutil.which("soffice")
    if not office:
        raise ValueError("旧版 .ppt 需要 LibreOffice；请先另存为 .pptx")
    profile = (output_dir / "libreoffice-profile").as_uri()
    result = subprocess.run(
        [office, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pptx",
         "--outdir", str(output_dir), str(path)],
        capture_output=True, text=True, check=False, timeout=120,
    )
    converted = output_dir / (path.stem + ".pptx")
    if result.returncode or not converted.is_file():
        raise ValueError(f"无法转换 .ppt：{result.stderr.strip() or result.stdout.strip()}")
    return converted


def extract(path: Path) -> tuple[str, list[str]]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return read_plain_text(path), []
    if suffix == ".pdf":
        return pdf_text(path)
    if suffix == ".pptx":
        return pptx_text(path)
    if suffix == ".docx":
        return docx_text(path)
    if suffix == ".ppt":
        with tempfile.TemporaryDirectory(prefix="course_ppt_") as tmp:
            return pptx_text(convert_legacy_ppt(path, Path(tmp)))
    raise ValueError(f"不支持的文件格式：{suffix}")


def import_one(source: Path, course: str, library: Path = LIBRARY) -> tuple[Path, list[str]]:
    source = source.expanduser().resolve()
    if not source.is_file():
        raise ValueError(f"找不到文件：{source}")
    if source.suffix.lower() not in SUPPORTED:
        raise ValueError(f"不支持的文件格式：{source.suffix}")
    body, warnings = extract(source)
    if len(body.strip()) < 10:
        raise ValueError(f"{source.name} 可提取的文字太少，暂不导入；请检查是否为扫描件或纯图片幻灯片")

    digest = file_digest(source)
    stem = safe_name(source.stem)
    folder = library / safe_name(course)
    originals = folder / "_originals"
    originals.mkdir(parents=True, exist_ok=True)
    archive = originals / f"{stem}_{digest}{source.suffix.lower()}"
    index = folder / f"{stem}_{digest}.md"
    if archive.resolve() != source:
        shutil.copy2(source, archive)
    index.write_text(
        f"# {source.name}\n\nSource file: `_originals/{archive.name}`\n\n{body.strip()}\n",
        encoding="utf-8",
    )
    return index, warnings


def main() -> None:
    parser = argparse.ArgumentParser(description="导入课程资料并生成可检索文字")
    parser.add_argument("--course", default="unsorted", help="课程目录名，例如 deep_learning")
    parser.add_argument("files", nargs="+", type=Path)
    args = parser.parse_args()
    failures = 0
    for source in args.files:
        try:
            index, warnings = import_one(source, args.course)
            print(f"已导入：{source.name} -> {index}")
            for warning in warnings:
                print(f"需检查：{warning}")
        except Exception as exc:
            print(f"导入失败：{source.name}：{exc}")
            failures += 1
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
